from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from datetime import date

def format_paragraph(p, text="", bold=False, size=11, align='left'):
    run = p.add_run(text)
    run.bold = bold
    p.style.font.name = "Arial"
    p.style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
   
    run.font.size = Pt(size)
    
    if align == 'center':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    elif align == 'right':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    else:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    return p

def gerar_contrato(nome_locador, nome_locatario, valor_aluguel, duracao_meses, endereco_imovel):
    doc = Document()

    #Cabeçalho
    section = doc.sections[0]
    header = section.header
    h_p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    h_p.clear()
    format_paragraph(h_p, "Lukas Rocha", bold=True, size=14, align='center')

    #Titulo
    p = doc.add_paragraph()
    format_paragraph(p, "Contrato de Locação Residencial", bold=True, size=12, align='center')
    doc.add_paragraph()    # linha em branco

    #Dados do contrato
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.autofit = True
    table.cell(0,0).paragraphs[0].add_run("Locador:").bold = True
    table.cell(0,1).paragraphs[0].add_run(nome_locador).bold = True
    table.cell(1,0).paragraphs[0].add_run("Locatário:").bold = True
    table.cell(1,1).paragraphs[0].add_run(nome_locatario).bold = True
    
    doc.add_paragraph() #Linha em branco

      # Informações do imóvel e valores
    p = doc.add_paragraph()
    format_paragraph(p, f'Imóvel: {endereco_imovel}', size=11)
    format_paragraph(doc.add_paragraph(), f'Valor do Aluguel: R$ {valor_aluguel} (reajustes conforme cláusula).', size=11)
    format_paragraph(doc.add_paragraph(), f'Duração: {duracao_meses} meses', size=11)
    doc.add_paragraph()

    # Cláusulas principais (exemplo completo)
    format_paragraph(doc.add_paragraph(), 'CLÁUSULAS', bold=True, size=12, align='center')
    doc.add_paragraph()

    clauses = [
        "1. PAGAMENTO: O locatário pagará o valor do aluguel até o dia 5 (cinco) de cada mês, via transferência bancária para conta indicada pelo locador.",
        "2. GARANTIA: As partes elegem a modalidade de garantia conforme acordo em anexo (ex: caução, fiador ou seguro fiança).",
        "3. MANUTENÇÃO: O locatário é responsável por pequenos reparos; o locador responde por problemas estruturais.",
        "4. RESCISÃO: Em caso de rescisão antecipada sem justa causa, a parte que rescindir pagará multa proporcional de 1 (um) mês por ano restante, proporcional ao período.",
        "5. UTILIDADES: Água, energia e demais taxas ordinárias são de responsabilidade do locatário, salvo disposição em contrário.",
        "6. VISTORIA: Será realizada vistoria de entrada e saída, assinalando estado do imóvel.",
        "7. FORO: Para dirimir dúvidas, fica eleito o foro da comarca do imóvel."
    ]

    for c in clauses:
        p = doc.add_paragraph()
        format_paragraph(p, c, size=11)

    doc.add_paragraph()
    # Cláusula de assinatura
    format_paragraph(doc.add_paragraph(), 'E por estarem justos e contratados, assinam o presente contrato em 2 (duas) vias de igual teor.', size=11)
    doc.add_paragraph()
    doc.add_paragraph()

    # Espaço para assinaturas
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.autofit = False
    sig_table.columns[0].width = Cm(8)
    sig_table.columns[1].width = Cm(8)
    sig_table.cell(0,0).paragraphs[0].add_run('__________________________').bold = False
    sig_table.cell(0,1).paragraphs[0].add_run('__________________________').bold = False
    sig_table.cell(1,0).paragraphs[0].add_run(f'Locador: {nome_locador}')
    sig_table.cell(1,1).paragraphs[0].add_run(f'Locatário: {nome_locatario}')

    # Rodapé com data
    today = date.today().strftime('%d/%m/%Y')
    footer_p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    footer_p.clear()
    format_paragraph(footer_p, f'Emitido em: {today} — Gerado por Sampaio Dev', size=9, align='center')

    # Salvar
    filename = f"Contrato_Aluguel_{nome_locatario.replace(' ', '_')}.docx"
    doc.save(filename)
    print(f'Contrato gerado: {filename}')


if __name__ == "__main__":
    nome_locador = input("Nome do Locador: ")
    nome_locatario = input("Nome do Locatário: ")
    valor_aluguel = input("Valor do Aluguel (R$): ")
    duracao_meses = input("Duração do Contrato (meses): ")
    endereco_imovel = input("Endereço do Imóvel: ")
    gerar_contrato( nome_locador, nome_locatario, valor_aluguel, duracao_meses, endereco_imovel)