# percorrer toda a nossa base de dados
    # para cada item, ver se o bairro ja existe uma aba, se não existir, criar aquela aba.
    # copiar os valores daquela linha e colocar na aba do bairro correspondente 

from pathlib import Path
from copy import copy

from openpyxl import load_workbook

def criar_aba(bairro, arquivo_base, estilo_cabecalho):
    if bairro not in arquivo_bairros.sheetnames:
        arquivo_bairros.create_sheet(bairro)
        nova_aba = arquivo_bairros[bairro]
        nova_aba["A1"].value = "Data de Nascimento"
        nova_aba["B1"].value = "Pessoa"
        nova_aba["C1"].value = "Bairro"
        nova_aba["A1"]._style = estilo_cabecalho
        nova_aba["B1"]._style = estilo_cabecalho
        nova_aba["C1"]._style = estilo_cabecalho

def transferir_informacoes_aba(aba_origem, aba_destino, linha_origem):
    linha_destino = aba_destino.max_row + 1
    for coluna in range(1, 4):
        celula_origem = aba_origem.cell(row=linha_origem, column=coluna)
        celula_destino = aba_destino.cell(row=linha_destino, column=coluna)
        celula_destino.value = celula_origem.value
        celula_destino._style = copy(celula_origem._style)

arquivo_bairros = load_workbook(Path(__file__).resolve().parent / "Bairros.xlsx")

print(arquivo_bairros.sheetnames)

aba_basedados = arquivo_bairros["Base de Dados"]

ultima_linha = aba_basedados.max_row
print(ultima_linha)

estilo_cabecalho = aba_basedados["A1"]._style

for linha in range(2, ultima_linha + 1):
    # bairro = aba_basedados.cell(row=linha, column=3).value
    bairro = aba_basedados[f"C{linha}"].value
    if not bairro:
        break
    # criar uma aba para o bairro

    criar_aba(bairro, arquivo_bairros, estilo_cabecalho)

    # transferir as informações para a aba correspondente

    aba_destino = arquivo_bairros[bairro]

    transferir_informacoes_aba(aba_basedados, aba_destino, linha)

arquivo_bairros.save("Bairros2.xlsx")

